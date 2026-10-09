"""
30 - QA checklist report (master, runs after 21): one PASS / WARN / FAIL / SKIP / HUMAN status per QA check
  Reads the JSON reports of the run (or the latest report of every script), decides every check of qa_checks.py
  from the issues those reports found, applies the people's results from qa_human_checks.json (qa.py mark), and
  writes Excel, JSON, CSV, Markdown (ready to paste into a ticket) and an HTML page to preview in the browser:
      verdict (NEEDS FIXES / REVIEW / READY), fixes required, warnings, skipped checks (and how to run them),
      human checks, notes for the dev team / questions for the client, the full checklist with evidence.

  python "py files/30_qa_checklist_report.py"
"""
import argparse
import html
import json
import os
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from qa_checks import BY_ID, CHECKS, SCRIPT_FOR_REPORT, SECTIONS
from seo_common import (CRITICAL, IMPORTANT, INFO, OPTIMIZATION, REPORT_ROOT, SCRIPTS_DIR, SEVERITY_ORDER, SITE_WIDE,
                        qa_config, rel_path, remove_old_reports, report_path, write_csv, write_sheet)

HUMAN_FILE = SCRIPTS_DIR / "qa_human_checks.json"
STATUS_ORDER = {"FAIL": 0, "WARN": 1, "SKIP": 2, "HUMAN": 3, "PASS": 4}
STATUS_FILL = {"FAIL": "FFC7CE", "WARN": "FFE4B5", "SKIP": "E7E6E6", "HUMAN": "E4DFEC", "PASS": "C6EFCE"}
EVIDENCE_MD, EVIDENCE_FULL = 8, 60


# ------------------------------------------------------------------ inputs
def load_reports():
    """{key: report JSON} for this run (SEO_RUN_TIME) or the latest of every script."""
    run_time, run_date = os.environ.get("SEO_RUN_TIME"), os.environ.get("SEO_RUN_DATE")
    found = {}
    for f in REPORT_ROOT.glob("*/json/*.json"):
        if f.name.startswith(("Website_Health_Report_", "QA_Checklist_Report_")):
            continue
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
            r = data["report"]
            key = r["key"]
        except (ValueError, KeyError, TypeError):
            continue
        stamp = (f.parent.parent.name, r.get("run_time", ""))
        if run_time and (stamp[1] != run_time or (run_date and stamp[0] != run_date)):
            continue
        if key not in found or stamp > found[key]["_stamp"]:
            data["_stamp"], data["_path"] = stamp, rel_path(f)
            found[key] = data
    return found


def load_human():
    """qa_human_checks.json (written by qa.py mark); an entry without a valid status is ignored with a warning."""
    try:
        data = json.loads(HUMAN_FILE.read_text(encoding="utf-8"))
    except OSError:
        return {}
    except ValueError as e:
        print(f"WARNING: {HUMAN_FILE.name} is not valid JSON ({e}) - the recorded human results are ignored")
        return {}
    if not isinstance(data, dict):
        print(f"WARNING: {HUMAN_FILE.name} must be a JSON object - the recorded human results are ignored")
        return {}
    out = {}
    for cid, entry in data.items():
        if str(cid).upper() not in BY_ID:
            print(f"WARNING: {HUMAN_FILE.name}: {cid} ignored (no such check - see qa.py list)")
        elif isinstance(entry, dict) and str(entry.get("status", "")).upper() in ("PASS", "FAIL", "WARN", "SKIP"):
            out[str(cid).upper()] = entry
        else:
            print(f"WARNING: {HUMAN_FILE.name}: {cid} ignored (needs a status pass / fail / warn / skip - use qa.py mark)")
    return out


# no options: -h shows this help instead of rebuilding the checklist, and an unknown option is an error
argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter).parse_args()
reports = load_reports()
human = load_human()
cfg = qa_config()
bases = Counter(r["report"].get("audited_server", "") for r in reports.values())
base = bases.most_common(1)[0][0] if bases else ""
issues_by_report = {}
for key, data in reports.items():
    items = [dict(i, url=u["url"]) for u in data.get("urls", []) for i in u["issues"]]
    items += [dict(i, url=SITE_WIDE) for i in data.get("site_wide_issues", [])]
    issues_by_report[key] = items
crawl = next((r["report"].get("crawl") for r in reports.values() if r["report"].get("crawl")), {}) or {}
# items a script could not check (an unexpected error on that page / URL): no check that reads that report may PASS
script_errors = {key: [i for i in items if i.get("category") == "Script error"] for key, items in issues_by_report.items()}
html_pages = crawl.get("html_pages") or 0


def coverage(key):
    """'' or a note when the report checked clearly fewer pages than the crawl found (e.g. --pages sample)."""
    r = reports[key]["report"]
    n = r.get("pages_checked") or 0
    if html_pages and n and n < 0.9 * html_pages and key not in ("28_repo",):
        return f"{key}: {n} of {html_pages} pages"
    return ""


# ------------------------------------------------------------------ evaluate
def evaluate(check):
    res = {"id": check.id, "section": check.section, "title": check.title, "kind": check.kind, "owner": check.owner,
           "how": check.how, "reports": check.reports, "auto_status": "", "status": "", "result": "", "evidence": [],
           "skipped": [], "partial": [], "manual": None, "counts": {}}
    if check.kind == "HUMAN":
        res["auto_status"] = "HUMAN"
    else:
        missing = [k for k in check.reports if k not in reports]
        matched, skipped = [], []
        for src in check.sources:
            for i in issues_by_report.get(src.report, []):
                if src.matches(i):
                    (skipped if i["issue"].startswith("Not checked:") else matched).append(dict(i, report_key=src.report))
        # the same issue can match two sources of a check: keep it once
        seen, uniq = set(), []
        for i in matched:
            k = (i["url"], i["issue"], i.get("current_value", ""), i.get("element", ""))
            if k not in seen:
                seen.add(k)
                uniq.append(i)
        matched = [i for i in uniq if i["severity"] != INFO]
        info = [i for i in uniq if i["severity"] == INFO]
        res["skipped"] = sorted({f"{i['issue'][13:]} ({i['current_value']})" for i in skipped})
        res["partial"] = [c for c in (coverage(k) for k in check.reports if k in reports) if c]
        counts = Counter(i["severity"] for i in matched)
        res["counts"] = {s: counts[s] for s in (CRITICAL, IMPORTANT, OPTIMIZATION)}
        matched.sort(key=lambda i: (SEVERITY_ORDER.get(i["severity"], 9), i["url"] == SITE_WIDE, i["url"]))
        res["evidence"] = matched + info
        if len(missing) == len(check.reports):
            res["auto_status"] = "SKIP"
            res["result"] = "not run: " + ", ".join(f"{SCRIPT_FOR_REPORT.get(k, k)} ({k})" for k in missing)
        elif matched:
            worst = min((i["severity"] for i in matched), key=lambda s: SEVERITY_ORDER[s])
            res["auto_status"] = "FAIL" if worst in check.fail else "WARN"
            pages = {i["url"] for i in matched if i["url"] != SITE_WIDE}
            names = Counter(i["issue"] for i in matched)
            top = "; ".join(f"{n} ({c}x)" if c > 1 else n for n, c in names.most_common(3))
            res["result"] = (f"{len(matched)} issue(s)" + (f" on {len(pages)} page(s)" if pages else "") + f": {top}"
                             + (f"; +{len(names) - 3} more issue types" if len(names) > 3 else ""))
        elif skipped and not matched:
            res["auto_status"] = "SKIP"
            res["result"] = "not checked: " + "; ".join(res["skipped"])[:300]
        else:
            res["auto_status"] = "PASS"
            res["result"] = "no issues found" + (f" ({len(info)} note(s))" if info else "")
        errors = [i for k in check.reports for i in script_errors.get(k, [])]
        if errors and res["auto_status"] == "PASS":
            res["auto_status"] = "SKIP"
            res["result"] = (f"not checked on {len(errors)} item(s) - the script failed there "
                             f"({errors[0]['current_value'][:120]})")
            res["evidence"] = errors + res["evidence"]
        elif errors and not any(i in res["evidence"] for i in errors):
            res["result"] += f" [+{len(errors)} item(s) not checked: script error]"
        if missing and res["auto_status"] != "SKIP":
            res["result"] += f" [not run: {', '.join(missing)}]"
        if res["partial"]:
            res["result"] += f" [partial run - {'; '.join(res['partial'])}]"
    res["status"] = res["auto_status"]
    h = human.get(check.id)
    if h and (not h.get("server") or not base or h.get("server") == base):
        res["manual"] = h
        res["status"] = h["status"].upper()
        res["result"] = f"{h['status'].upper()} by {h.get('by') or 'QA'} on {h.get('date', '')}: {h.get('note', '')}" + \
                        (f" (automatic: {res['auto_status']} - {res['result']})" if check.kind != "HUMAN" else "")
    elif check.kind == "AUTO+HUMAN" and res["auto_status"] in ("PASS", "WARN"):
        res["result"] += " - also confirm by hand: " + check.how if check.how else ""
    return res


results = [evaluate(c) for c in CHECKS]
by_status = defaultdict(list)
for r in results:
    by_status[r["status"]].append(r)
counts = {s: len(by_status[s]) for s in STATUS_ORDER}
pending_human = [r for r in results if r["status"] == "HUMAN"]
if counts["FAIL"]:
    verdict = "NEEDS FIXES"
elif counts["WARN"] or counts["SKIP"] or pending_human:
    verdict = "REVIEW"
else:
    verdict = "READY"
now = datetime.now()
report_rows = [(k, d["report"]["title"], f"{d['_stamp'][0]} {d['_stamp'][1]}", d["report"].get("pages_checked", ""),
                d["report"].get("score", ""), d["report"].get("audited_server", ""), d["_path"])
               for k, d in sorted(reports.items())]
mixed = len(bases) > 1
stale = sorted({d["_stamp"][0] for d in reports.values()})


def ev_line(i, limit=160):
    url = i["url"].replace(base, "") or "/" if i["url"].startswith("http") else i["url"]
    cur = (i.get("current_value") or "")[:limit]
    el = (i.get("element") or "")[:90]
    return f"{url}: [{i['severity']}] {i['issue']}" + (f" - {cur}" if cur else "") + (f" ({el})" if el and el not in cur else "")


# ------------------------------------------------------------------ Markdown
md = [f"# QA Report - {cfg['company'] or base}", ""]
md += [f"- **Verdict:** **{verdict}**", f"- **Generated:** {now:%Y-%m-%d %H:%M}", f"- **Audited server:** {base}"]
if cfg["live_url"]:
    md.append(f"- **Live site:** {cfg['live_url']}")
if cfg["phones"]:
    md.append(f"- **Confirmed phone:** {', '.join(cfg['phones'])}")
md.append(f"- **Legal site:** {'yes' if cfg['legal_site'] else 'no'}")
md.append(f"- **Reports used:** {len(reports)} ({', '.join(sorted(reports))})")
md.append(f"- **Checks:** {len(results)} total - FAIL {counts['FAIL']}, WARN {counts['WARN']}, PASS {counts['PASS']}, "
          f"SKIP {counts['SKIP']}, HUMAN {counts['HUMAN']}")
if crawl:
    md.append(f"- **Pages:** {crawl.get('html_pages')} HTML pages crawled ({crawl.get('sitemap_urls')} in the sitemap)")
if mixed:
    md.append(f"- **Warning:** the reports come from different servers: {dict(bases)} - re-run them against one server")
if len(stale) > 1:
    md.append(f"- **Note:** reports from several dates are combined ({', '.join(stale)}) - re-run old ones for a fresh result")
md.append("")


def md_section(title, items, with_evidence=True):
    if not items:
        return
    md.extend([f"## {title}", ""])
    for n, r in enumerate(items, 1):
        md.append(f"{n}. **{r['id']}** {r['title']} - {r['result']} _(owner: {r['owner']})_")
        if with_evidence:
            for i in r["evidence"][:EVIDENCE_MD]:
                md.append(f"   - `{ev_line(i)}`")
            if len(r["evidence"]) > EVIDENCE_MD:
                md.append(f"   - ...and {len(r['evidence']) - EVIDENCE_MD} more (see the Excel / HTML report)")
    md.append("")


md_section("Fixes required", by_status["FAIL"])
md_section("Warnings to review", by_status["WARN"])
if by_status["SKIP"]:
    md.extend(["## Not checked (SKIP)", ""])
    for r in by_status["SKIP"]:
        md.append(f"- **{r['id']}** {r['title']} - {r['result']}" + (f" - how: {r['how']}" if r["how"] else ""))
    md.append("")
md.extend(["## Human checks", ""])
for r in results:
    if r["kind"] != "AUTO":
        done = r["manual"] and r["manual"]["status"].upper() in ("PASS",)
        state = f" - {r['manual']['status'].upper()}: {r['manual'].get('note', '')}" if r["manual"] else ""
        md.append(f"- [{'x' if done else ' '}] **{r['id']}** {r['title']}{state}" + (f" _(how: {r['how']})_" if r["how"]
                                                                                       and not r["manual"] else ""))
md.append("")


def note_lines(owners):
    out = []
    for r in by_status["FAIL"] + by_status["WARN"]:
        if r["owner"] in owners:
            first = r["evidence"][0] if r["evidence"] else None
            out.append(f"- {r['id']} {r['title']}: {r['result']}" + (f" e.g. {ev_line(first, 100)}" if first else ""))
    return out


client_q = [f"- {r['id']} {r['title']}: {r['result']}" for r in by_status["WARN"] + by_status["FAIL"]
            if r["owner"] == "client"] + [f"- {r['id']} {r['title']} ({r['how']})" for r in pending_human
                                          if r["owner"] == "client"]
for title, lines in (("Notes for the dev team", note_lines({"dev"})),
                     ("Notes for SEO / content", note_lines({"seo", "content"})),
                     ("Questions for the client", client_q)):
    if lines:
        md.extend([f"## {title}", "", "```", *lines, "```", ""])
md.extend(["## Full checklist", ""])
for section in SECTIONS:
    md.extend([f"### {section}", "", "| ID | Check | Type | Status | Result | Owner |", "|----|-------|------|--------|--------|-------|"])
    for r in results:
        if r["section"] == section:
            cell = r["result"].replace("|", "\\|").replace("\n", " ")
            md.append(f"| {r['id']} | {r['title']} | {r['kind']} | {r['status']} | {cell[:220]} | {r['owner']} |")
    md.append("")
md.extend(["## Reports used", "", "| Key | Report | Run | Pages | Score | File |", "|---|---|---|---|---|---|"])
md += [f"| {k} | {t} | {run} | {p} | {s} | `{path}` |" for k, t, run, p, s, _, path in report_rows]
md.append("")

# ------------------------------------------------------------------ HTML
E = html.escape


def pill(s):
    return f'<span class="pill {s.lower()}">{s}</span>'


cards = "".join(f'<button class="card {s.lower()}" data-filter="{s}"><b>{counts[s]}</b>{s}</button>' for s in STATUS_ORDER)
rows_html = []
for section in SECTIONS:
    rows_html.append(f'<section><h2>{E(section)}</h2>')
    for r in results:
        if r["section"] != section:
            continue
        ev = "".join(f"<li>{E(ev_line(i, 300))}</li>" for i in r["evidence"][:EVIDENCE_FULL])
        if len(r["evidence"]) > EVIDENCE_FULL:
            ev += f"<li>…and {len(r['evidence']) - EVIDENCE_FULL} more in the Excel report</li>"
        extra = (f'<p class="how">How: {E(r["how"])}</p>' if r["how"] else "") + \
                (f'<p class="how">Reports: {E(", ".join(r["reports"]))}</p>' if r["reports"] else "")
        body = f"<ul>{ev}</ul>" if ev else ""
        rows_html.append(
            f'<details class="check" data-status="{r["status"]}" data-text="{E((r["id"] + " " + r["title"] + " " + r["result"]).lower())}">'
            f'<summary>{pill(r["status"])}<code>{r["id"]}</code><span class="t">{E(r["title"])}</span>'
            f'<span class="k">{r["kind"]} · {r["owner"]}</span><span class="r">{E(r["result"])}</span></summary>'
            f'{extra}{body}</details>')
    rows_html.append('</section>')
page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>QA Checklist</title><style>
:root{{--bg:#f6f7fb;--fg:#1d2433;--card:#fff;--muted:#5d6678;--line:#e3e6ee;
--fail:#e5484d;--warn:#f5a524;--pass:#30a46c;--skip:#8b8d98;--human:#8e4ec6}}
@media (prefers-color-scheme:dark){{:root{{--bg:#111318;--fg:#e8eaf0;--card:#1a1d24;--muted:#9aa1b2;--line:#2a2f3a}}}}
*{{box-sizing:border-box}}body{{margin:0;background:var(--bg);color:var(--fg);font:15px/1.45 system-ui,-apple-system,Segoe UI,sans-serif}}
main{{max-width:1180px;margin:0 auto;padding:24px 16px 60px}}h1{{margin:0 0 4px;font-size:26px}}
.meta{{color:var(--muted);margin:0 0 18px}}.verdict{{display:inline-block;padding:6px 14px;border-radius:999px;font-weight:700;color:#fff;
background:{'var(--fail)' if verdict == 'NEEDS FIXES' else 'var(--warn)' if verdict == 'REVIEW' else 'var(--pass)'}}}
.cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(120px,1fr));gap:10px;margin:18px 0}}
.card{{border:0;border-radius:14px;padding:12px;color:#fff;font:600 13px system-ui;cursor:pointer;text-align:left}}
.card b{{display:block;font-size:28px}}.card.fail{{background:var(--fail)}}.card.warn{{background:var(--warn)}}
.card.pass{{background:var(--pass)}}.card.skip{{background:var(--skip)}}.card.human{{background:var(--human)}}
.tools{{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:8px}}.tools input{{flex:1;min-width:200px;padding:9px 12px;border-radius:10px;
border:1px solid var(--line);background:var(--card);color:var(--fg)}}.tools button{{padding:8px 12px;border-radius:10px;border:1px solid var(--line);
background:var(--card);color:var(--fg);cursor:pointer}}h2{{font-size:17px;margin:26px 0 8px}}
.check{{background:var(--card);border:1px solid var(--line);border-radius:12px;margin:6px 0;padding:0 12px}}
summary{{display:grid;grid-template-columns:70px 78px 1fr 150px;gap:10px;align-items:center;padding:10px 0;cursor:pointer;list-style:none}}
summary>*{{min-width:0;overflow-wrap:anywhere}}summary .r{{grid-column:3/5;color:var(--muted);font-size:13px}}summary .k{{color:var(--muted);font-size:12px;text-align:right}}
.pill{{border-radius:999px;padding:2px 0;text-align:center;font:700 12px system-ui;color:#fff}}.pill.fail{{background:var(--fail)}}
.pill.warn{{background:var(--warn)}}.pill.pass{{background:var(--pass)}}.pill.skip{{background:var(--skip)}}.pill.human{{background:var(--human)}}
code{{font-weight:700}}.how{{color:var(--muted);margin:2px 0 8px;font-size:13px}}ul{{margin:0 0 12px;padding-left:18px;font:12.5px ui-monospace,Menlo,monospace;
overflow-wrap:anywhere}}pre{{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px;white-space:pre-wrap;overflow-wrap:anywhere;font-size:12.5px}}
@media (max-width:700px){{summary{{grid-template-columns:62px 1fr}}summary .t{{grid-column:1/3}}summary .k{{display:none}}summary .r{{grid-column:1/3}}}}
</style></head><body><main>
<h1>QA Checklist</h1><p class="meta">{E(cfg['company'] or base)} · {now:%Y-%m-%d %H:%M} · {len(reports)} reports · {E(base)}</p>
<span class="verdict">{verdict}</span>
{'<p class="meta">⚠ Reports from different servers are mixed: ' + E(str(dict(bases))) + '</p>' if mixed else ''}
<div class="cards">{cards}</div>
<div class="tools"><input id="q" placeholder="Search checks, issues, URLs…"><button data-filter="ALL">Show all</button>
<button id="open">Expand all</button></div>
{''.join(rows_html)}
<h2>Notes for the dev team</h2><pre>{E(chr(10).join(note_lines({"dev"})) or 'none')}</pre>
<h2>Notes for SEO / content</h2><pre>{E(chr(10).join(note_lines({"seo", "content"})) or 'none')}</pre>
<h2>Questions for the client</h2><pre>{E(chr(10).join(client_q) or 'none')}</pre>
</main><script>
let filter='ALL';const q=document.getElementById('q');
function apply(){{const t=q.value.toLowerCase();document.querySelectorAll('.check').forEach(c=>{{
c.style.display=(filter==='ALL'||c.dataset.status===filter)&&(!t||c.dataset.text.includes(t)||c.innerText.toLowerCase().includes(t))?'':'none';}});
document.querySelectorAll('section').forEach(s=>{{s.style.display=[...s.querySelectorAll('.check')].some(c=>c.style.display!=='none')?'':'none';}});}}
document.querySelectorAll('[data-filter]').forEach(b=>b.onclick=()=>{{filter=b.dataset.filter;apply();}});q.oninput=apply;
document.getElementById('open').onclick=()=>document.querySelectorAll('.check').forEach(c=>{{if(c.style.display!=='none')c.open=true;}});
</script></body></html>"""

# ------------------------------------------------------------------ Excel / JSON / CSV
name = "QA_Checklist_Report"
remove_old_reports(name)
xlsx, jpath, cpath, mpath, hpath = (report_path(name, e) for e in ("xlsx", "json", "csv", "md", "html"))
wb = Workbook()
ws = wb.active
ws.title = "Summary"
for row in (["QA Checklist Report"], ["Verdict", verdict], ["Generated", now.strftime("%Y-%m-%d %H:%M:%S")],
            ["Audited server", base], *[[s, counts[s]] for s in STATUS_ORDER], ["Reports used", len(reports)]):
    ws.append(row)
ws["A1"].font = Font(bold=True, size=16)
ws["B2"].fill = PatternFill(fill_type="solid", start_color=STATUS_FILL["FAIL" if verdict == "NEEDS FIXES" else
                                                                       "WARN" if verdict == "REVIEW" else "PASS"])
for i, s in enumerate(STATUS_ORDER, 5):
    ws.cell(row=i, column=1).fill = PatternFill(fill_type="solid", start_color=STATUS_FILL[s])
ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 60
headers = ["ID", "Section", "Check", "Type", "Status", "Automatic status", "Result", "Owner", "Critical", "Important",
           "Optimization", "How / notes", "Reports"]
rows = [(r["id"], r["section"], r["title"], r["kind"], r["status"], r["auto_status"], r["result"], r["owner"],
         r["counts"].get(CRITICAL, ""), r["counts"].get(IMPORTANT, ""), r["counts"].get(OPTIMIZATION, ""), r["how"],
         ", ".join(r["reports"])) for r in sorted(results, key=lambda r: (STATUS_ORDER[r["status"]], r["id"]))]
cws = wb.create_sheet("Checklist")
write_sheet(cws, headers, rows, (9, 22, 60, 12, 9, 12, 90, 9, 9, 10, 12, 50, 30))
for n in range(2, cws.max_row + 1):
    for col in (5, 6):
        v = cws.cell(row=n, column=col).value
        if v in STATUS_FILL:
            cws.cell(row=n, column=col).fill = PatternFill(fill_type="solid", start_color=STATUS_FILL[v])
    cws.cell(row=n, column=7).alignment = Alignment(wrap_text=True, vertical="top")
ev_rows = [(r["id"], r["status"], i["severity"], i["url"], i["issue"], i.get("current_value", ""), i.get("element", ""),
            i.get("recommended_fix", ""), i.get("report", "")) for r in results for i in r["evidence"][:500]]
write_sheet(wb.create_sheet("Evidence"), ["Check", "Status", "Severity", "URL", "Issue", "Current value", "Element",
                                          "Recommended fix", "Report"], ev_rows, (9, 8, 13, 55, 45, 50, 40, 60, 24),
            severity_col=3)
write_sheet(wb.create_sheet("Human checks"), ["ID", "Check", "Status", "How", "Recorded"],
            [(r["id"], r["title"], r["status"], r["how"],
              json.dumps(r["manual"]) if r["manual"] else "") for r in results if r["kind"] != "AUTO"], (9, 60, 9, 70, 50))
write_sheet(wb.create_sheet("Reports"), ["Key", "Report", "Run", "Pages", "Score", "Server", "JSON"], report_rows,
            (20, 30, 20, 8, 8, 30, 60))
wb.save(xlsx)
jpath.write_text(json.dumps({
    "report": {"title": "QA Checklist Report", "verdict": verdict, "generated": now.isoformat(timespec="seconds"),
               "audited_server": base, "counts": counts, "checks": len(results), "reports_used": sorted(reports),
               "mixed_servers": dict(bases) if mixed else {}, "excel": rel_path(xlsx), "markdown": rel_path(mpath),
               "html": rel_path(hpath)},
    "checks": [{k: v for k, v in r.items() if k != "evidence"} | {"evidence": r["evidence"][:200]} for r in results],
}, indent=1, ensure_ascii=False, default=str))
write_csv(cpath, headers, rows)
mpath.write_text("\n".join(md), encoding="utf-8")
hpath.write_text(page, encoding="utf-8")

print("\n" + "=" * 70)
print("✓ QA Checklist Master Report generated successfully!")
print("=" * 70)
print(f"  • Excel Report:    {rel_path(xlsx)}")
print(f"  • JSON Report:     {rel_path(jpath)}")
print(f"  • CSV Report:      {rel_path(cpath)}")
print(f"  • Markdown Report: {rel_path(mpath)}")
print(f"  • HTML Report:     {rel_path(hpath)}")
print("-" * 70)
print(f"  Verdict: {verdict} | " + "  ".join(f"{s}: {counts[s]}" for s in STATUS_ORDER))
if by_status["FAIL"]:
    print(f"  Failing Checks ({len(by_status['FAIL'])}):")
    for r in by_status["FAIL"][:8]:
        print(f"    ✗ {r['id']:<8} {r['title'][:65]}")
    if len(by_status["FAIL"]) > 8:
        print(f"    ... and {len(by_status['FAIL']) - 8} more failures (see HTML/Excel report)")
print("=" * 70 + "\n")
